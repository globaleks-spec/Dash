"""Dump selected workbook ranges with formula and cached values."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


def serialize(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("sheet")
    parser.add_argument("ranges", nargs="+")
    args = parser.parse_args()

    formula_book = load_workbook(args.path, data_only=False, read_only=False)
    value_book = load_workbook(args.path, data_only=True, read_only=False)
    formula_sheet = formula_book[args.sheet]
    value_sheet = value_book[args.sheet]

    result: dict[str, Any] = {
        "sheet": args.sheet,
        "max_row": formula_sheet.max_row,
        "max_column": formula_sheet.max_column,
        "ranges": [],
    }
    for requested_range in args.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(requested_range)
        rows: list[dict[str, Any]] = []
        for row in range(min_row, max_row + 1):
            cells: list[dict[str, Any]] = []
            for col in range(min_col, max_col + 1):
                formula_cell = formula_sheet.cell(row=row, column=col)
                value_cell = value_sheet.cell(row=row, column=col)
                if formula_cell.value is None and value_cell.value is None:
                    continue
                item: dict[str, Any] = {
                    "cell": formula_cell.coordinate,
                    "value": serialize(value_cell.value),
                }
                if formula_cell.data_type == "f":
                    item["formula"] = serialize(formula_cell.value)
                elif formula_cell.value != value_cell.value:
                    item["source"] = serialize(formula_cell.value)
                item["format"] = formula_cell.number_format
                cells.append(item)
            if cells:
                rows.append({"row": row, "cells": cells})
        result["ranges"].append({"range": requested_range, "rows": rows})

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
