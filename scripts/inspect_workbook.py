"""Read-only structural inspection for weekly sales workbooks."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def inspect_workbook(path: Path, max_rows: int, max_cols: int) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    result: dict[str, Any] = {
        "path": str(path),
        "sheet_names": workbook.sheetnames,
        "sheets": [],
    }

    for sheet in workbook.worksheets:
        rows: list[dict[str, Any]] = []
        for row_index in range(1, min(sheet.max_row, max_rows) + 1):
            cells: list[dict[str, Any]] = []
            for col_index in range(1, min(sheet.max_column, max_cols) + 1):
                cell = sheet.cell(row=row_index, column=col_index)
                if cell.value is not None:
                    cells.append(
                        {
                            "coordinate": cell.coordinate,
                            "value": json_value(cell.value),
                            "data_type": cell.data_type,
                            "number_format": cell.number_format,
                        }
                    )
            if cells:
                rows.append({"row": row_index, "cells": cells})

        result["sheets"].append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "hidden": sheet.sheet_state != "visible",
                "hidden_rows": [
                    index for index, dim in sheet.row_dimensions.items() if dim.hidden
                ],
                "hidden_columns": [
                    key for key, dim in sheet.column_dimensions.items() if dim.hidden
                ],
                "sample_rows": rows,
            }
        )
    return result


def compact_profile(path: Path, max_rows: int, max_cols: int) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    payload: dict[str, Any] = {"path": str(path), "sheets": []}
    for sheet in workbook.worksheets:
        row_profiles: list[dict[str, Any]] = []
        limit_rows = min(sheet.max_row, max_rows)
        limit_cols = min(sheet.max_column, max_cols)
        for row_index in range(1, limit_rows + 1):
            values: list[dict[str, Any]] = []
            string_count = 0
            numeric_count = 0
            formula_count = 0
            for col_index in range(1, limit_cols + 1):
                cell = sheet.cell(row=row_index, column=col_index)
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    formula_count += 1
                elif isinstance(cell.value, str):
                    string_count += 1
                elif isinstance(cell.value, (int, float)):
                    numeric_count += 1
                rendered = str(json_value(cell.value)).replace("\n", " ")
                values.append(
                    {
                        "cell": cell.coordinate,
                        "value": rendered[:160],
                    }
                )
            if values:
                row_profiles.append(
                    {
                        "row": row_index,
                        "non_empty": len(values),
                        "strings": string_count,
                        "numbers": numeric_count,
                        "formulas": formula_count,
                        "values": values[:30],
                    }
                )

        header_candidates = sorted(
            row_profiles,
            key=lambda item: (
                item["strings"] * 3 + item["non_empty"] - item["formulas"] * 2,
                item["non_empty"],
            ),
            reverse=True,
        )[:4]
        payload["sheets"].append(
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "merged_count": len(sheet.merged_cells.ranges),
                "merged_ranges": [str(item) for item in list(sheet.merged_cells.ranges)[:20]],
                "header_candidates": header_candidates,
            }
        )
    return payload


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("--max-rows", type=int, default=40)
    parser.add_argument("--max-cols", type=int, default=30)
    parser.add_argument("--compact", action="store_true")
    args = parser.parse_args()

    if args.compact:
        payload = compact_profile(args.path, args.max_rows, args.max_cols)
    else:
        payload = inspect_workbook(args.path, args.max_rows, args.max_cols)
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
