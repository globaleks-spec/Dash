from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def make_weekly_workbook(path: Path, *, changed: bool = False) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Реализация БП 2026"
    sheet["C1"] = "Стоимость указана в т.р. без НДС"
    sheet["A3"] = "Направление/ Менеджер"
    sheet["B3"] = "Контрагент"
    sheet["AM3"] = "БП 2026"

    groups = [
        ("AN", "AS", "1 квартал 2026", ["БП 1 квартал", "Факт отгрузки", "Планируемая отгрузка", "Подписанные контракты", "Проекты акцепт Заказчика", "Прочие проекты в проработке"]),
        ("AT", "AZ", "2 квартал 2026", ["БП 2 квартал План", "БП-2 + ГЭП", "Факт отгрузки", "Планируемая отгрузка", "Подписанные контракты", "Проекты акцепт Заказчика", "Прочие проекты в проработке"]),
        ("BD", "BJ", "3 квартал 2026", ["БП 3 квартал План", "БП-3 + ГЭП", "Факт отгрузки", "Планируемая отгрузка", "Подписанные контракты", "Проекты акцепт Заказчика", "Прочие проекты в проработке"]),
        ("BN", "BT", "4 квартал 2026", ["БП 4 квартал План", "БП-4 + ГЭП", "Факт отгрузки", "Планируемая отгрузка", "Подписанные контракты", "Проекты акцепт Заказчика", "Прочие проекты в проработке"]),
        ("BX", "CD", "1 квартал 2027", ["БП 1 квартал", "Факт отгрузки", "Планируемая отгрузка", "Подписанные контракты", "Проекты акцепт Заказчика", "Прочие проекты в проработке", "Прочие проекты в проработке"]),
    ]
    from openpyxl.utils import column_index_from_string

    for start, end, label, headers in groups:
        sheet.merge_cells(f"{start}2:{end}2")
        sheet[f"{start}2"] = label
        start_col = column_index_from_string(start)
        for offset, header in enumerate(headers):
            sheet.cell(3, start_col + offset).value = header

    sheet["A4"] = "СЗГД"
    sheet["B4"] = "ИТОГО"
    sheet["AM4"] = 1000
    sheet["AN4"] = 250
    sheet["AT4"] = 250
    sheet["BD4"] = 250
    sheet["BN4"] = 250

    sheet["A5"] = "Иванов И.И."
    sheet["B5"] = "ИТОГО"
    sheet["AM5"] = 1000
    sheet["AN5"] = 250
    sheet["AT5"] = 250
    sheet["BD5"] = 250
    sheet["BN5"] = 250

    sheet["A6"] = "Иванов И.И."
    sheet["B6"] = "Заказчик А ИБП" + (" - отгружен 15.10.26" if changed else "")
    sheet["AO6"] = 100

    sheet["A7"] = "Иванов И.И."
    sheet["B7"] = "Заказчик Б СЭЩ ШОТ"
    sheet["BJ7"] = 300
    sheet["BT7"] = 300

    sheet["A8"] = "Иванов И.И."
    sheet["B8"] = "Дублируемый проект СЭЩ сч.№77-2026"
    sheet["BR8"] = 50
    sheet["A9"] = "Иванов И.И."
    sheet["B9"] = "Дублируемый проект СЭЩ сч.№77-2026"
    sheet["BR9"] = 50

    sheet["A10"] = "Иванов И.И."
    sheet["B10"] = "Проект следующего года СОПТ"
    sheet["CC10"] = 75

    data_sheet = workbook.create_sheet("Данные_new")
    data_sheet.append(["Менеджер СКД", "План", "Факт", "Прогноз", "% факт", "Дефицит", "Риск"])
    workbook.save(path)
    return path
