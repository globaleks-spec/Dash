from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, BinaryIO, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .db import initialize, transaction


MAIN_SHEET_MARKERS = {"Направление/ Менеджер", "Контрагент"}
OUTSTANDING_METRICS = ("confirmed", "signed", "accepted", "pipeline")
METRIC_PRIORITY = {"confirmed": 4, "signed": 3, "accepted": 2, "pipeline": 1}


@dataclass(slots=True)
class ImportResult:
    import_id: int
    filename: str
    projects: int
    managers: int
    raw_cells: int
    duplicate_groups: int
    already_imported: bool = False


class MappingRequired(ValueError):
    pass


def _bytes(source: bytes | bytearray | str | Path | BinaryIO) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    data = source.read()
    if hasattr(source, "seek"):
        source.seek(0)
    return data


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _text(value)).casefold().replace("ё", "е")
    return re.sub(r"\s+", " ", text).strip()


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None
    candidate = value.replace("\xa0", "").replace(" ", "").replace(",", ".")
    candidate = re.sub(r"[^0-9.\-]", "", candidate)
    if not candidate or candidate in {"-", ".", "-."}:
        return None
    try:
        return float(candidate)
    except ValueError:
        return None


def parse_week_from_filename(filename: str) -> tuple[str | None, str | None]:
    match = re.search(
        r"(\d{1,2}\.\d{1,2}\.(?:\d{2}|\d{4}))\s*[-–—]\s*"
        r"(\d{1,2}\.\d{1,2}\.(?:\d{2}|\d{4}))",
        filename,
    )
    if not match:
        return None, None

    def parse(raw: str) -> str:
        fmt = "%d.%m.%Y" if len(raw.rsplit(".", 1)[-1]) == 4 else "%d.%m.%y"
        return datetime.strptime(raw, fmt).date().isoformat()

    return parse(match.group(1)), parse(match.group(2))


def _header_candidates(sheet: Any, limit: int = 40) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for row in range(1, min(sheet.max_row, limit) + 1):
        values = [_text(sheet.cell(row, col).value) for col in range(1, sheet.max_column + 1)]
        populated = [(index + 1, value) for index, value in enumerate(values) if value]
        if len(populated) < 2:
            continue
        strings = sum(not value.replace(".", "", 1).isdigit() for _, value in populated)
        candidates.append(
            {
                "row": row,
                "score": strings * 3 + len(populated),
                "columns": [
                    {"column": get_column_letter(col), "name": value[:160]}
                    for col, value in populated[:80]
                ],
            }
        )
    return sorted(candidates, key=lambda item: item["score"], reverse=True)[:5]


def profile_workbook(source: bytes | str | Path | BinaryIO) -> dict[str, Any]:
    data = _bytes(source)
    workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        candidates = _header_candidates(sheet)
        sheets.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "header_candidates": candidates,
            }
        )
    primary = detect_main_sheet(workbook)
    mapping: list[dict[str, Any]] = []
    if primary:
        sheet = workbook[primary]
        mapping.extend(
            [
                {"source": "A", "header": _text(sheet["A3"].value), "target": "manager", "confidence": "высокая"},
                {"source": "B", "header": _text(sheet["B3"].value), "target": "project_name + source_text", "confidence": "высокая"},
            ]
        )
        for item in detect_metric_columns(sheet):
            mapping.append(
                {
                    "source": item["letter"],
                    "header": item["header"],
                    "target": f"{item['year']} Q{item['quarter']} · {item['metric']}",
                    "confidence": "высокая",
                }
            )
    return {"sheets": sheets, "primary_sheet": primary, "suggested_mapping": mapping}


def detect_main_sheet(workbook: Any) -> str | None:
    for sheet in workbook.worksheets:
        headers = {_text(sheet.cell(3, col).value) for col in range(1, min(sheet.max_column, 150) + 1)}
        if MAIN_SHEET_MARKERS.issubset(headers):
            return sheet.title
    return None


def _merged_row_labels(sheet: Any, row: int) -> dict[int, str]:
    labels = {col: _text(sheet.cell(row, col).value) for col in range(1, sheet.max_column + 1)}
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row:
            label = _text(sheet.cell(merged.min_row, merged.min_col).value)
            for col in range(merged.min_col, merged.max_col + 1):
                labels[col] = label
    return labels


def _metric_from_header(header: str) -> str | None:
    value = _normalized(header)
    if "факт отгруз" in value:
        return "actual"
    if "планируемая отгруз" in value:
        return "confirmed"
    if "подписан" in value and "контракт" in value:
        return "signed"
    if "акцепт" in value:
        return "accepted"
    if "проч" in value and "проект" in value:
        return "pipeline"
    if "бп" in value and "гэп" in value:
        return "adjusted_plan"
    if "бп" in value and "план" in value:
        return "plan"
    return None


def detect_metric_columns(sheet: Any) -> list[dict[str, Any]]:
    row_two = _merged_row_labels(sheet, 2)
    result: list[dict[str, Any]] = []
    for col in range(1, sheet.max_column + 1):
        group = row_two.get(col, "")
        match = re.search(r"([1-4])\s*квартал(?:\s*(20\d{2}))?", _normalized(group))
        if not match or not match.group(2):
            continue
        metric = _metric_from_header(_text(sheet.cell(3, col).value))
        if not metric:
            continue
        result.append(
            {
                "column": col,
                "letter": get_column_letter(col),
                "year": int(match.group(2)),
                "quarter": int(match.group(1)),
                "metric": metric,
                "header": _text(sheet.cell(3, col).value),
                "group": group,
            }
        )
    return result


def _annual_plan_columns(sheet: Any) -> dict[int, int]:
    result: dict[int, int] = {}
    for col in range(1, sheet.max_column + 1):
        match = re.fullmatch(r"бп\s*(20\d{2})", _normalized(sheet.cell(3, col).value))
        if match:
            result[int(match.group(1))] = col
    return result


def _is_total(value: Any) -> bool:
    return _normalized(value) in {"итого", "итог"}


def _scope(name: str, project_managers: set[str]) -> str:
    normalized = _normalized(name)
    if normalized == "сзгд":
        return "company"
    if name in project_managers:
        return "manager"
    if normalized.startswith("оп ") or "отдел" in normalized:
        return "department"
    return "other"


STATUS_TAIL = re.compile(
    r"\s[-–—]\s(?:проект\b|прогноз\b|контракт\b|отгруж\w*\b|перенос\b|"
    r"поставк\w*\b|первая очередь\b|вторая очередь\b).*$",
    flags=re.IGNORECASE,
)


def _business_base(text: str) -> str:
    base = STATUS_TAIL.sub("", text).strip()
    base = re.sub(r"\(\s*отгруж\w*[^)]*\)", "", base, flags=re.IGNORECASE)
    base = re.sub(r"\b(?:ао\s+)?сэщ\b", "", base, flags=re.IGNORECASE)
    base = re.sub(r"\bао\s*[«\"]?гк\s*[«\"]?электрощит[^()\-–—,;]*", "", base, flags=re.IGNORECASE)
    base = _normalized(base)
    base = re.sub(r"[^0-9a-zа-я№]+", " ", base)
    return re.sub(r"\s+", " ", base).strip()


def business_key(manager: str, source_text: str) -> str:
    identifiers = re.findall(
        r"(?:сч\.?\s*№?|р\.?\s*№?|закупк\w*\s*№?|заказ\w*\s*№?)\s*"
        r"([0-9а-яa-z\-/]+)",
        _normalized(source_text),
    )
    identity = "|".join(sorted(set(identifiers))) or _business_base(source_text)
    payload = f"{_normalized(manager)}|{identity}"
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()[:24]


def _extract_date_near(text: str, keywords: Iterable[str]) -> str | None:
    keyword_pattern = "|".join(re.escape(item) for item in keywords)
    pattern = re.compile(
        rf"(?:{keyword_pattern})[^0-9]{{0,35}}(\d{{1,2}}\.\d{{1,2}}\.(?:\d{{2}}|\d{{4}}))",
        flags=re.IGNORECASE,
    )
    matches = pattern.findall(text)
    parsed: list[date] = []
    for raw in matches:
        fmt = "%d.%m.%Y" if len(raw.rsplit(".", 1)[-1]) == 4 else "%d.%m.%y"
        try:
            parsed.append(datetime.strptime(raw, fmt).date())
        except ValueError:
            continue
    return max(parsed).isoformat() if parsed else None


def _channel(text: str) -> str | None:
    value = _normalized(text)
    if "сэщ" in value or "электрощит" in value:
        return "СЭЩ"
    if re.search(r"\bкит\b", value):
        return "КИТ"
    if "агент" in value:
        return "Агент"
    return None


def _equipment(text: str) -> str | None:
    patterns = [
        (r"\bибп\b", "ИБП"),
        (r"\bсбп\b", "СБП"),
        (r"\bзв[уy]\b", "ЗВУ"),
        (r"\bшот\b", "ШОТ"),
        (r"\bшуот\b", "ШУОТ"),
        (r"\bсопт\b", "СОПТ"),
        (r"выпрямител", "Выпрямители"),
        (r"аккумулятор|\bакб\b", "Аккумуляторные батареи"),
        (r"стабилиз", "Стабилизаторы"),
    ]
    lowered = _normalized(text)
    found = [label for pattern, label in patterns if re.search(pattern, lowered)]
    return ", ".join(dict.fromkeys(found)) or None


def _final_customer(text: str, channel: str | None) -> str | None:
    if channel != "СЭЩ":
        return None
    match = re.match(r"(.{3,120}?)\s+(?:ао\s+)?сэщ\b", text, flags=re.IGNORECASE)
    return match.group(1).strip(" -–—,;") if match else None


def _stage(text: str, current_metric: str | None, actual_amount: float) -> str | None:
    value = _normalized(text)
    if "отмен" in value:
        return "отменено"
    if "не состоится" in value or "проигран" in value or "проиграли" in value:
        return "проиграно"
    if "перенес" in value or "сдвиг" in value:
        return "перенесено"
    if actual_amount > 0 or "отгружен" in value or "отгружена" in value:
        return "отгружено"
    if current_metric == "signed" or ("договор" in value and "подпис" in value):
        return "подписанный договор"
    if current_metric == "accepted" or "акцепт" in value:
        return "акцепт"
    if "изготовлен" in value:
        return "изготовлено"
    if "производств" in value:
        return "производство"
    if "тендер" in value or "конкурс" in value:
        return "участие в тендере"
    if "техническ" in value and "оцен" in value:
        return "техническая оценка"
    if "ткп" in value:
        return "ТКП"
    if "техническ" in value or "проработк" in value:
        return "техническая проработка"
    return None


def _risk_flags(text: str) -> dict[str, int]:
    value = _normalized(text)
    return {
        "risk_financial": int(
            any(term in value for term in ("нет финансирования", "бюджет не утвержден", "нет бюджета", "нет лимит"))
        ),
        "risk_technical": int(
            any(term in value for term in ("техническ", "проектирован", "согласование техники"))
        ),
        "risk_production": int(
            any(term in value for term in ("производственн", "изготовлен", "в производстве"))
        ),
        "risk_procurement": int(
            any(term in value for term in ("комплектующ", "закупочн", "снабжен"))
        ),
        "risk_schedule": int(
            any(term in value for term in ("перенос", "перенес", "сдвиг", "задерж", "просроч"))
        ),
    }


def _customer(text: str) -> str:
    value = STATUS_TAIL.sub("", text).strip()
    return value[:250]


def _current_status(
    values: list[dict[str, Any]], year: int
) -> tuple[str | None, int | None, float, bool]:
    outstanding = [
        item
        for item in values
        if item["year"] == year
        and item["metric"] in OUTSTANDING_METRICS
        and abs(item["amount"]) > 1e-9
    ]
    if not outstanding:
        return None, None, 0.0, False
    latest_quarter = max(item["quarter"] for item in outstanding)
    latest = [item for item in outstanding if item["quarter"] == latest_quarter]
    conflict = len({item["metric"] for item in latest}) > 1
    selected = max(
        latest,
        key=lambda item: (METRIC_PRIORITY[item["metric"]], abs(item["amount"])),
    )
    return selected["metric"], latest_quarter, float(selected["amount"]), conflict


def _year_from_week(week_end: str | None, metric_columns: list[dict[str, Any]]) -> int:
    if week_end:
        return date.fromisoformat(week_end).year
    years = Counter(item["year"] for item in metric_columns)
    return years.most_common(1)[0][0] if years else datetime.now().year


def _insert_raw_cells(
    connection: sqlite3.Connection,
    import_id: int,
    formula_book: Any,
    value_book: Any,
) -> int:
    rows: list[tuple[Any, ...]] = []
    count = 0
    for sheet in formula_book.worksheets:
        cached_sheet = value_book[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                cached = cached_sheet.cell(cell.row, cell.column).value
                if cell.value is None and cached is None:
                    continue
                rows.append(
                    (
                        import_id,
                        sheet.title,
                        cell.row,
                        cell.column,
                        cell.coordinate,
                        json.dumps(_json_value(cell.value), ensure_ascii=False),
                        json.dumps(_json_value(cached), ensure_ascii=False),
                        cell.data_type,
                        cell.number_format,
                    )
                )
                count += 1
                if len(rows) >= 5000:
                    connection.executemany(
                        "INSERT INTO raw_cells(import_id, sheet_name, row_index, column_index, "
                        "coordinate, raw_value, cached_value, data_type, number_format) "
                        "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        rows,
                    )
                    rows.clear()
    if rows:
        connection.executemany(
            "INSERT INTO raw_cells(import_id, sheet_name, row_index, column_index, "
            "coordinate, raw_value, cached_value, data_type, number_format) "
            "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)",
            rows,
        )
    return count


def _insert_sheet_metadata(connection: sqlite3.Connection, import_id: int, workbook: Any) -> None:
    for sheet in workbook.worksheets:
        candidates = _header_candidates(sheet)
        connection.execute(
            "INSERT INTO sheets(import_id, name, state, max_row, max_column, header_row) "
            "VALUES(?, ?, ?, ?, ?, ?)",
            (
                import_id,
                sheet.title,
                sheet.sheet_state,
                sheet.max_row,
                sheet.max_column,
                candidates[0]["row"] if candidates else None,
            ),
        )


def import_workbook(
    connection: sqlite3.Connection,
    source: bytes | str | Path | BinaryIO,
    filename: str | None = None,
) -> ImportResult:
    initialize(connection)
    data = _bytes(source)
    filename = filename or (Path(source).name if isinstance(source, (str, Path)) else "upload.xlsx")
    digest = hashlib.sha256(data).hexdigest()
    existing = connection.execute(
        "SELECT id FROM imports WHERE file_hash = ?", (digest,)
    ).fetchone()
    if existing:
        project_count = connection.execute(
            "SELECT COUNT(*) FROM project_snapshots WHERE import_id = ?", (existing["id"],)
        ).fetchone()[0]
        manager_count = connection.execute(
            "SELECT COUNT(DISTINCT manager) FROM project_snapshots WHERE import_id = ?", (existing["id"],)
        ).fetchone()[0]
        raw_count = connection.execute(
            "SELECT COUNT(*) FROM raw_cells WHERE import_id = ?", (existing["id"],)
        ).fetchone()[0]
        duplicate_count = connection.execute(
            "SELECT COUNT(DISTINCT duplicate_group) FROM project_snapshots "
            "WHERE import_id = ? AND is_ambiguous = 1",
            (existing["id"],),
        ).fetchone()[0]
        return ImportResult(
            int(existing["id"]), filename, project_count, manager_count, raw_count, duplicate_count, True
        )

    formula_book = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    value_book = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    primary_sheet = detect_main_sheet(formula_book)
    if not primary_sheet:
        raise MappingRequired(
            "Не найден лист с заголовками «Направление/ Менеджер» и «Контрагент». "
            "Выберите лист и соответствие полей вручную."
        )

    week_start, week_end = parse_week_from_filename(filename)
    sheet = value_book[primary_sheet]
    formula_sheet = formula_book[primary_sheet]
    unit_note = " ".join(_text(sheet.cell(1, col).value) for col in range(1, min(60, sheet.max_column) + 1))
    multiplier = 1000.0 if re.search(r"\bт\.?\s*р", _normalized(unit_note)) else 1.0
    source_unit = "тыс. руб. без НДС" if multiplier == 1000 else "руб. без НДС"
    metric_columns = detect_metric_columns(sheet)
    active_year = _year_from_week(week_end, metric_columns)
    annual_columns = _annual_plan_columns(sheet)

    meta = {
        "sheet_names": formula_book.sheetnames,
        "detected_mapping": metric_columns,
        "active_year": active_year,
        "source_note": unit_note,
    }

    with transaction(connection):
        cursor = connection.execute(
            "INSERT INTO imports(file_hash, filename, week_start, week_end, imported_at, "
            "primary_sheet, source_unit, unit_multiplier, status, workbook_meta_json) "
            "VALUES(?, ?, ?, ?, ?, ?, ?, ?, 'completed', ?)",
            (
                digest,
                filename,
                week_start,
                week_end,
                datetime.now().isoformat(timespec="seconds"),
                primary_sheet,
                source_unit,
                multiplier,
                json.dumps(meta, ensure_ascii=False),
            ),
        )
        import_id = int(cursor.lastrowid)
        _insert_sheet_metadata(connection, import_id, formula_book)
        raw_count = _insert_raw_cells(connection, import_id, formula_book, value_book)

        project_rows = [
            row
            for row in range(4, sheet.max_row + 1)
            if _text(sheet.cell(row, 1).value)
            and _text(sheet.cell(row, 2).value)
            and not _is_total(sheet.cell(row, 2).value)
        ]
        project_managers = {_text(sheet.cell(row, 1).value) for row in project_rows}

        for row in range(4, sheet.max_row + 1):
            if not _is_total(sheet.cell(row, 2).value):
                continue
            name = _text(sheet.cell(row, 1).value)
            if not name:
                continue
            scope = _scope(name, project_managers)
            stored_name = "__TOTAL__" if scope == "company" else name
            for year, col in annual_columns.items():
                amount = _number(sheet.cell(row, col).value)
                if amount is not None:
                    connection.execute(
                        "INSERT OR IGNORE INTO manager_plans(import_id, manager, scope, year, quarter, "
                        "metric, amount, source_sheet, source_row, source_cell) "
                        "VALUES(?, ?, ?, ?, NULL, 'annual_plan', ?, ?, ?, ?)",
                        (
                            import_id,
                            stored_name,
                            scope,
                            year,
                            amount * multiplier,
                            primary_sheet,
                            row,
                            f"{get_column_letter(col)}{row}",
                        ),
                    )
            for column in metric_columns:
                if column["metric"] not in {"plan", "adjusted_plan"}:
                    continue
                amount = _number(sheet.cell(row, column["column"]).value)
                if amount is None:
                    continue
                connection.execute(
                    "INSERT OR IGNORE INTO manager_plans(import_id, manager, scope, year, quarter, "
                    "metric, amount, source_sheet, source_row, source_cell) "
                    "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        import_id,
                        stored_name,
                        scope,
                        column["year"],
                        column["quarter"],
                        column["metric"],
                        amount * multiplier,
                        primary_sheet,
                        row,
                        f"{column['letter']}{row}",
                    ),
                )

        snapshot_rows: list[dict[str, Any]] = []
        conflict_keys: set[str] = set()
        for row in project_rows:
            manager = _text(sheet.cell(row, 1).value)
            source_text = _text(sheet.cell(row, 2).value)
            key = business_key(manager, source_text)
            values: list[dict[str, Any]] = []
            source_values: dict[str, Any] = {}
            for column in metric_columns:
                raw_amount = _number(sheet.cell(row, column["column"]).value)
                if raw_amount is None or abs(raw_amount) < 1e-9:
                    continue
                amount = raw_amount * multiplier
                item = {**column, "amount": amount, "raw_amount": raw_amount}
                values.append(item)
                source_values[f"{column['letter']}{row}"] = raw_amount
                connection.execute(
                    "INSERT INTO period_values(import_id, business_key, source_sheet, source_row, "
                    "year, quarter, metric, amount, source_cell, raw_amount) "
                    "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        import_id,
                        key,
                        primary_sheet,
                        row,
                        column["year"],
                        column["quarter"],
                        column["metric"],
                        amount,
                        f"{column['letter']}{row}",
                        raw_amount,
                    ),
                )

            current_metric, planned_quarter, outstanding_amount, conflict = _current_status(values, active_year)
            actual_amount = sum(
                item["amount"]
                for item in values
                if item["year"] == active_year and item["metric"] == "actual"
            )
            next_year_amounts = [
                item for item in values if item["year"] > active_year and item["metric"] in OUTSTANDING_METRICS
            ]
            planned_year = active_year if current_metric else None
            if not current_metric and next_year_amounts:
                planned_year = min(item["year"] for item in next_year_amounts)
                next_values = [item for item in next_year_amounts if item["year"] == planned_year]
                current_metric, planned_quarter, outstanding_amount, next_conflict = _current_status(
                    next_values, planned_year
                )
                conflict = conflict or next_conflict
            if conflict:
                conflict_keys.add(key)

            channel = _channel(source_text)
            stage = _stage(source_text, current_metric, actual_amount)
            risks = _risk_flags(source_text)
            record = {
                "import_id": import_id,
                "business_key": key,
                "source_sheet": primary_sheet,
                "source_row": row,
                "manager": manager,
                "customer": _customer(source_text),
                "final_customer": _final_customer(source_text, channel),
                "project_name": source_text,
                "channel": channel,
                "industry": None,
                "equipment_type": _equipment(source_text),
                "project_amount": actual_amount + outstanding_amount if actual_amount + outstanding_amount else None,
                "stage": stage,
                "contract_date": _extract_date_near(source_text, ("договор", "контракт")),
                "planned_ship_date": _extract_date_near(source_text, ("плановая дата отгрузки",)),
                "actual_ship_date": _extract_date_near(source_text, ("отгружен", "отгружена", "отгружено")),
                "planned_year": planned_year,
                "planned_quarter": planned_quarter,
                "probability": None,
                **risks,
                "next_owner": None,
                "next_action": None,
                "next_action_due": None,
                "comment": None,
                "current_category": current_metric,
                "source_text": source_text,
                "source_values_json": json.dumps(source_values, ensure_ascii=False),
                "last_updated": week_end or datetime.now().date().isoformat(),
                "duplicate_group": None,
                "is_ambiguous": int(conflict),
            }
            snapshot_rows.append(record)

        groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for record in snapshot_rows:
            groups[record["business_key"]].append(record)
        duplicate_groups = 0
        for key, records in groups.items():
            if len(records) > 1:
                duplicate_groups += 1
                for record in records:
                    record["duplicate_group"] = key
                    record["is_ambiguous"] = 1
                connection.execute(
                    "INSERT OR IGNORE INTO review_queue(import_id, business_key, issue_type, details, "
                    "created_at) VALUES(?, ?, 'ambiguous_duplicate', ?, ?)",
                    (
                        import_id,
                        key,
                        f"Найдено строк с одинаковым ключом проекта: {len(records)}. Требуется проверка.",
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )

        fields = list(snapshot_rows[0]) if snapshot_rows else []
        if fields:
            placeholders = ", ".join("?" for _ in fields)
            connection.executemany(
                f"INSERT INTO project_snapshots({', '.join(fields)}) VALUES({placeholders})",
                [[record[field] for field in fields] for record in snapshot_rows],
            )

        for record in snapshot_rows:
            if not record["project_amount"]:
                connection.execute(
                    "INSERT OR IGNORE INTO review_queue(import_id, business_key, issue_type, details, "
                    "created_at) VALUES(?, ?, 'missing_amount', ?, ?)",
                    (
                        import_id,
                        record["business_key"],
                        f"{record['source_sheet']}!B{record['source_row']}: отсутствует сумма проекта.",
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )
            if record["business_key"] in conflict_keys:
                connection.execute(
                    "INSERT OR IGNORE INTO review_queue(import_id, business_key, issue_type, details, "
                    "created_at) VALUES(?, ?, 'conflicting_status', ?, ?)",
                    (
                        import_id,
                        record["business_key"],
                        f"{record['source_sheet']}!{record['source_row']}: в одном квартале несколько стадий.",
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )

    managers = len({record["manager"] for record in snapshot_rows if record["manager"]})
    return ImportResult(import_id, filename, len(snapshot_rows), managers, raw_count, duplicate_groups)


def result_as_dict(result: ImportResult) -> dict[str, Any]:
    return asdict(result)
